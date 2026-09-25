"""
Script para generar la documentación completa en formato Word (.docx)
explicando detalladamente el uso del sistema, la formulación matemática de los 4 modelos
de predicción y el algoritmo de reabastecimiento en OperaStock.
"""
import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls


def set_cell_background(cell, fill_hex):
    """Establece el color de fondo de una celda en Word."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Establece los márgenes internos de una celda."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)


def add_callout(doc, text, title="💡 NOTA EXPLICATIVA PARA LA TESIS"):
    """Agrega una caja destacada (Callout Box) al documento Word."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F0Fdf4") # Verde suave
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Borde izquierdo verde
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="16A34A"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"{title}\n")
    run_t.bold = True
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(22, 163, 74)

    run_b = p.add_run(text)
    run_b.font.name = 'Calibri'
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph() # Espacio


def generar_word(filename):
    doc = Document()

    # Configuración de márgenes a 2.5 cm (1 pulgada)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Estilos globales
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(51, 65, 85) # Slate

    # ═════════════════════════════════════════════════════════════════════════
    # PORTADA / ENCABEZADO PRINCIPAL
    # ═════════════════════════════════════════════════════════════════════════
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_sub = p_title.add_run("PROYECTO DE TESIS — DOCUMENTACIÓN TÉCNICA Y METODOLÓGICA\n")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(11)
    r_sub.bold = True
    r_sub.font.color.rgb = RGBColor(37, 99, 235) # Azul primary

    r_main = p_title.add_run("GUÍA EXPLICATIVA DEL MÓDULO DE REABASTECIMIENTO Y MODELOS PREDICTIVOS")
    r_main.font.name = 'Calibri'
    r_main.font.size = Pt(20)
    r_main.bold = True
    r_main.font.color.rgb = RGBColor(15, 23, 42) # Dark Navy

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(24)
    r_meta = p_meta.add_run("Sistema OperaStock · Control Predictivo de Compras e Inventario\nMoneda de Costeo: Soles (S/) · Cobertura Geográfica: Perú")
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_heading("1. Resumen Ejecutivo y Flujo General del Sistema", level=1)
    
    p = doc.add_paragraph(
        "El módulo de abastecimiento de OperaStock fue diseñado para resolver el problema crítico "
        "de gestión de inventarios en plantas de producción: determinar el momento exacto (¿Cuándo comprar?) "
        "y la cantidad óptima (¿Cuánto comprar?) de materias primas sin incurrir en paradas de planta por desabastecimiento "
        "ni en sobrecostos por exceso de almacenamiento."
    )
    p.paragraph_format.space_after = Pt(10)

    doc.add_heading("1.1. Flujo Metodológico en 7 Pasos", level=2)

    steps = [
        ("1. Registro de Consumos Reales", "El sistema captura diariamente el insumo consumido en cada producción confirmada."),
        ("2. Filtrado de Anomalías (IQR)", "Se aíslan valores atípicos mediante el Rango Intercuartílico para evitar sesgar las predicciones."),
        ("3. Competencia Multimodelo en Paralelo", "Se entrenan 4 modelos (Random Forest, SARIMA, Holt-Winters y Promedio Móvil) sobre el 80% de los datos."),
        ("4. Evaluación y Selección Óptima", "Se valida el desempeño de los 4 modelos sobre el 20% restante (Holdout Chronological) y se selecciona el de menor RMSE."),
        ("5. Cálculo de Stock de Seguridad (SS)", "Se determina el colchón de reserva dinámico en función de la desviación estándar y el tiempo de entrega del proveedor."),
        ("6. Alerta de Punto de Reorden (ROP)", "Se compara el Stock Utilizable actual contra el ROP. Si Stock ≤ ROP, se dispara la sugerencia de compra."),
        ("7. Ajuste Comercial y Gobernanza", "Se aplica el MOQ y múltiplos del proveedor en Soles (S/). El usuario puede editar la cantidad antes de aprobar la orden de compra.")
    ]

    for title, desc in steps:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(f"{title}: ")
        r1.bold = True
        r1.font.color.rgb = RGBColor(15, 23, 42)
        r2 = p.add_run(desc)

    add_callout(
        doc,
        "Esta arquitectura garantiza que la planta nunca dependa de la intuición o de cálculos estáticos. "
        "Si los consumos se vuelven estables, el modelo ajusta el stock a la baja para ahorrar capital; si hay picos, "
        "el algoritmo reacciona ampliando la cobertura de protección.",
        title="📌 APORTE DE TESIS"
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SECCIÓN 2: EXPLICACIÓN DETALLADA DE LOS 4 MODELOS PREDICTIVOS
    # ═════════════════════════════════════════════════════════════════════════
    doc.add_heading("2. Explicación Detallada de los 4 Modelos Predictivos", level=1)
    
    doc.add_paragraph(
        "A diferencia de los software comerciales tradicionales que aplican una fórmula fija para todos los insumos, "
        "OperaStock ejecuta un esquema competitivo de 4 algoritmos adaptativos:"
    )

    # MODELO 1: RANDOM FOREST
    doc.add_heading("2.1. Random Forest Regressor (Machine Learning)", level=2)
    doc.add_paragraph(
        "Random Forest es un algoritmo de aprendizaje supervisado basado en un ensamble de múltiples árboles de decisión. "
        "Para predecir el consumo futuro, el modelo genera cientos de árboles independientes entrenados con subconjuntos aleatorios de datos "
        "y promedia sus resultados para reducir la variancia."
    )
    
    # Tabla de características Random Forest
    rf_table = doc.add_table(rows=4, cols=2)
    rf_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rf_table.style = 'Table Grid'
    
    rf_data = [
        ("Ecuación / Lógica de Ensamble", "y_hat = (1 / B) * Σ T_b(x)   [Promedio de las predicciones de B árboles]"),
        ("Features (Variables de Entrada)", "Consumos pasados en t-7, t-14, t-21, t-30 días, promedios móviles y desviación estándar."),
        ("Parámetros en OperaStock", "n_estimators=100, max_depth=8, min_samples_leaf=2, random_state=42."),
        ("¿Por qué se utiliza en la tesis?", "Captura relaciones complejas no lineales que los modelos estadísticos lineales no pueden detectar.")
    ]
    
    for idx, (k, v) in enumerate(rf_data):
        cell_k = rf_table.cell(idx, 0)
        cell_v = rf_table.cell(idx, 1)
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.3)
        set_cell_background(cell_k, "F1F5F9")
        set_cell_margins(cell_k, 80, 80, 100, 100)
        set_cell_margins(cell_v, 80, 80, 100, 100)
        
        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_after = Pt(2)
        rk = pk.add_run(k)
        rk.bold = True
        rk.font.size = Pt(9.5)
        
        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_after = Pt(2)
        rv = pv.add_run(v)
        rv.font.size = Pt(9.5)

    doc.add_paragraph() # Espacio

    # MODELO 2: SARIMA
    doc.add_heading("2.2. SARIMA (Seasonal AutoRegressive Integrated Moving Average)", level=2)
    doc.add_paragraph(
        "SARIMA es el modelo estadístico por excelencia para series de tiempo temporales con comportamiento cíclico. "
        "Combina componentes autorregresivos (AR), diferenciación de integración (I) para volver la serie estacionaria, "
        "y componentes de promedio móvil (MA), tanto a nivel ordinario como estacional."
    )

    sarima_table = doc.add_table(rows=4, cols=2)
    sarima_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sarima_table.style = 'Table Grid'

    sarima_data = [
        ("Ecuación General", "Φ_P(B^s) * φ_p(B) * (1-B)^d * (1-B^s)^D * y_t = Θ_Q(B^s) * θ_q(B) * ε_t"),
        ("Parámetros Ordinarios (p, d, q)", "order = (1, 1, 1) -> 1 término autorregresivo, 1 diferenciación, 1 término MA."),
        ("Parámetros Estacionales (P, D, Q)_s", "seasonal_order = (1, 0, 1, 7) -> Período estacional de 7 días (semanal)."),
        ("¿Por qué se utiliza en la tesis?", "Modela con precisión patrones estacionales repetitivos (ej. aumentos de consumo los viernes).")
    ]

    for idx, (k, v) in enumerate(sarima_data):
        cell_k = sarima_table.cell(idx, 0)
        cell_v = sarima_table.cell(idx, 1)
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.3)
        set_cell_background(cell_k, "F1F5F9")
        set_cell_margins(cell_k, 80, 80, 100, 100)
        set_cell_margins(cell_v, 80, 80, 100, 100)

        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_after = Pt(2)
        rk = pk.add_run(k)
        rk.bold = True
        rk.font.size = Pt(9.5)

        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_after = Pt(2)
        rv = pv.add_run(v)
        rv.font.size = Pt(9.5)

    doc.add_paragraph()

    # MODELO 3: HOLT-WINTERS
    doc.add_heading("2.3. Holt-Winters (Suavizado Exponencial Triple)", level=2)
    doc.add_paragraph(
        "El modelo de Holt-Winters descompone la demanda histórica en tres componentes dinámicos: "
        "el nivel base (ℓ_t), la tendencia (b_t) y la estacionalidad (s_t). Aplica factores de suavizado exponencial "
        "que le otorgan mayor peso a los datos más recientes."
    )

    hw_table = doc.add_table(rows=4, cols=2)
    hw_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hw_table.style = 'Table Grid'

    hw_data = [
        ("Ecuaciones de Suavizado", "Nivel: ℓ_t = α(y_t - s_{t-m}) + (1-α)(ℓ_{t-1} + b_{t-1})\nTendencia: b_t = β(ℓ_t - ℓ_{t-1}) + (1-β)b_{t-1}\nEstacionalidad: s_t = γ(y_t - ℓ_{t-1} - b_{t-1}) + (1-γ)s_{t-m}"),
        ("Ecuación de Pronóstico", "y_hat_{t+h} = ℓ_t + h * b_t + s_{t+h-m(k+1)}"),
        ("Parámetros en OperaStock", "trend = 'add', seasonal = 'add', seasonal_periods = 7."),
        ("¿Por qué se utiliza en la tesis?", "Reacciona inmediatamente si la planta acelera su ritmo de consumo diario de una semana a otra.")
    ]

    for idx, (k, v) in enumerate(hw_data):
        cell_k = hw_table.cell(idx, 0)
        cell_v = hw_table.cell(idx, 1)
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.3)
        set_cell_background(cell_k, "F1F5F9")
        set_cell_margins(cell_k, 80, 80, 100, 100)
        set_cell_margins(cell_v, 80, 80, 100, 100)

        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_after = Pt(2)
        rk = pk.add_run(k)
        rk.bold = True
        rk.font.size = Pt(9.5)

        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_after = Pt(2)
        rv = pv.add_run(v)
        rv.font.size = Pt(9.5)

    doc.add_paragraph()

    # MODELO 4: PROMEDIO MÓVIL
    doc.add_heading("2.4. Promedio Móvil (Baseline / Referencia Mínima)", level=2)
    doc.add_paragraph(
        "El Promedio Móvil calcula la media aritmética simple sobre los últimos N días de producción. "
        "Es un modelo determinista que asume consumo futuro uniforme."
    )

    pm_table = doc.add_table(rows=4, cols=2)
    pm_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    pm_table.style = 'Table Grid'

    pm_data = [
        ("Ecuación Matemática", "d_promedio = (1 / N) * Σ y_{t-i+1}   [para i=1 hasta N]"),
        ("Configuración en OperaStock", "Ventana móvil N = 7 días de operación."),
        ("Función 1: Baseline de Control", "Permite validar científicamente si los modelos avanzados de Machine Learning aportan una mejora real frente a la media."),
        ("Función 2: Fallback de Seguridad", "Si una materia prima es nueva (menos de 30 datos) o los modelos avanzados no convergen, el sistema cae de manera segura a esta fórmula para no detener la operación.")
    ]

    for idx, (k, v) in enumerate(pm_data):
        cell_k = pm_table.cell(idx, 0)
        cell_v = pm_table.cell(idx, 1)
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.3)
        set_cell_background(cell_k, "F1F5F9")
        set_cell_margins(cell_k, 80, 80, 100, 100)
        set_cell_margins(cell_v, 80, 80, 100, 100)

        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_after = Pt(2)
        rk = pk.add_run(k)
        rk.bold = True
        rk.font.size = Pt(9.5)

        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_after = Pt(2)
        rv = pv.add_run(v)
        rv.font.size = Pt(9.5)

    doc.add_paragraph()

    # ═════════════════════════════════════════════════════════════════════════
    # SECCIÓN 3: MÉTRICAS Y SELECCIÓN DEL MODELO GANADOR
    # ═════════════════════════════════════════════════════════════════════════
    doc.add_heading("3. Métricas de Evaluación y Selección del Modelo Óptimo", level=1)
    
    doc.add_paragraph(
        "Para decidir de manera imparcial cuál de los 4 modelos debe generar las compras, "
        "OperaStock utiliza la validación cruzada cronológica (*Holdout Chronological Validation*) reservando el último 20% del historial. "
        "Se calculan dos métricas principales:"
    )

    doc.add_paragraph(
        "1. Error Medio Absoluto (MAE):\n"
        "   MAE = (1 / n) * Σ |y_i - y_hat_i|\n"
        "   Mide la magnitud promedio de los errores en las mismas unidades de la materia prima (ej. kilos errados)."
    )
    doc.add_paragraph(
        "2. Raíz del Error Cuadrático Medio (RMSE):\n"
        "   RMSE = √[ (1 / n) * Σ (y_i - y_hat_i)² ]\n"
        "   Penaliza con mayor severidad los errores grandes o desviaciones atípicas."
    )

    add_callout(
        doc,
        "Criterio de Selección: El sistema compara el RMSE de los 4 competidores. "
        "El modelo con MENOR RMSE es guardado en formato binario (.joblib) con su hash SHA-256 para garantizar integridad auditada.",
        title="🎯 CRITERIO DE SELECCIÓN DE LA IA"
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SECCIÓN 4: FORMULACIÓN MATEMÁTICA DEL REABASTECIMIENTO (ROP Y SS)
    # ═════════════════════════════════════════════════════════════════════════
    doc.add_heading("4. Formulación Matemática del Reabastecimiento", level=1)

    doc.add_paragraph(
        "Una vez obtenido el consumo diario proyectado (d_hat) del modelo ganador, "
        "se aplican las ecuaciones de la teoría de control de inventarios:"
    )

    doc.add_heading("4.1. Stock de Seguridad Dinámico (SS)", level=2)
    doc.add_paragraph("SS = Z * σ_d * √(Lead Time)")
    doc.add_paragraph(
        "• Z = Nivel de servicio deseado (Z = 1.645 para 95% de confianza contra agotamiento).\n"
        "• σ_d = Desviación estándar del consumo diario en la ventana de evaluación.\n"
        "• Lead Time = Tiempo de entrega en días del proveedor asignado en Perú."
    )

    doc.add_heading("4.2. Punto de Reorden (ROP)", level=2)
    doc.add_paragraph("ROP = (d_hat * Lead Time) + SS")
    doc.add_paragraph(
        "El ROP establece el umbral de activación de la orden de compra. "
        "Si Stock Utilizable ≤ ROP, el sistema genera inmediatamente la sugerencia de reabastecimiento."
    )

    doc.add_heading("4.3. Lote Óptimo de Compra (Q) y Ajuste Comercial", level=2)
    doc.add_paragraph("Q_bruto = (d_hat * Cobertura Días) + SS - Stock Utilizable - Pedidos en Tránsito")
    doc.add_paragraph("Q_final = max( MOQ, ceil( Q_bruto / Múltiplo ) * Múltiplo )")
    doc.add_paragraph(
        "Esta formulación garantiza que la recomendación matemática respete los mínimos de venta (MOQ) "
        "y múltiplos de presentación del proveedor (ej. sacos de 5 kg o galones)."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SECCIÓN 5: GOBERNANZA HUMAN-IN-THE-LOOP Y AUDITORÍA
    # ═════════════════════════════════════════════════════════════════════════
    doc.add_heading("5. Gobernanza Human-in-the-Loop y Auditoría", level=1)
    
    doc.add_paragraph(
        "OperaStock implementa la filosofía Human-in-the-Loop. El algoritmo propone la cantidad matemáticamente óptima (ej. 55 kg), "
        "pero el comprador humano puede ajustar la cantidad a pedir (ej. 65 kg) antes de emitir la Orden de Compra."
    )
    doc.add_paragraph(
        "Cada modificación manual queda registrada en el módulo de auditoría con la firma del usuario, "
        "guardando la discrepancia entre la IA y la decisión humana para evaluaciones posteriores en la tesis."
    )

    # Guardar documento
    doc.save(filename)
    print(f"Documento Word creado exitosamente en: {filename}")


if __name__ == '__main__':
    output_docx = os.path.join(os.getcwd(), 'scripts', 'Guia_Explicativa_Modelos_Predictivos_OperaStock.docx')
    generar_word(output_docx)
